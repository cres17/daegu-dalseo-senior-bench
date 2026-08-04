#!/usr/bin/env python3
"""Validate newly added source files and compare them with the prior profile.

This script is read-only with respect to project source files. It writes audit
tables under ``02_data_understanding/profile_outputs``.
"""

from __future__ import annotations

import importlib.util
import json
import unicodedata
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "02_data_understanding" / "profile_outputs"
BASELINE = ROOT / "02_data_understanding" / "profile_outputs_baseline_20260722"


def load_module(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


profiler = load_module("profiler", Path(__file__).with_name("profile_sources.py"))
core = load_module("core_checks", Path(__file__).with_name("analyze_key_relationships.py"))


def nfc(value: object) -> str:
    return unicodedata.normalize("NFC", str(value)).strip()


def find_path(token: str, suffix: str) -> Path:
    matches = [
        path
        for path in ROOT.rglob(f"*{suffix}")
        if "02_data_understanding" not in path.parts and token in nfc(str(path))
    ]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one {suffix} path containing {token!r}, got {matches}")
    return matches[0]


def read_source_csv(token: str) -> pd.DataFrame:
    return core.read_csv(find_path(token, ".csv"))


def metric(rows: list[dict[str, object]], dataset: str, name: str, value: object, note: str = "") -> None:
    rows.append({"dataset": dataset, "metric": name, "value": value, "note": note})


def inventory_changes() -> pd.DataFrame:
    old = pd.read_csv(BASELINE / "file_inventory.csv")
    new = pd.read_csv(OUT / "file_inventory.csv")
    merged = old.merge(new, on="path", how="outer", suffixes=("_old", "_new"), indicator=True)
    records = []
    for _, row in merged.iterrows():
        if row["_merge"] == "left_only":
            change = "removed"
        elif row["_merge"] == "right_only":
            change = "added"
        elif row["sha256_old"] != row["sha256_new"]:
            change = "changed"
        else:
            change = "unchanged"
        records.append(
            {
                "path": row["path"],
                "change": change,
                "rows_old": row.get("rows_old", ""),
                "rows_new": row.get("rows_new", ""),
                "columns_old": row.get("columns_old", ""),
                "columns_new": row.get("columns_new", ""),
                "sha256_old": row.get("sha256_old", ""),
                "sha256_new": row.get("sha256_new", ""),
            }
        )
    return pd.DataFrame(records)


def latest_code_sheet_profile(xlsx_path: Path) -> tuple[pd.DataFrame, dict[str, object]]:
    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    sheet = workbook["2025년 6월"]
    rows = list(sheet.iter_rows(min_row=3, values_only=True))
    columns = ["시도코드", "시도명칭", "시군구코드", "시군구명칭", "읍면동코드", "읍면동명칭"]
    frame = pd.DataFrame(rows, columns=columns).dropna(how="all")
    dalseo = frame[(frame["시도명칭"] == "대구광역시") & (frame["시군구명칭"] == "달서구")].copy()
    dalseo["ADM_CD"] = (
        dalseo["시도코드"].astype(str).str.replace(r"\.0$", "", regex=True).str.zfill(2)
        + dalseo["시군구코드"].astype(str).str.replace(r"\.0$", "", regex=True).str.zfill(3)
        + dalseo["읍면동코드"].astype(str).str.replace(r"\.0$", "", regex=True).str.zfill(3)
    )
    profile = {
        "workbook_sheets": len(workbook.sheetnames),
        "latest_sheet": sheet.title,
        "latest_sheet_data_rows": len(frame),
        "dalseo_rows": len(dalseo),
    }
    return dalseo, profile


def main() -> None:
    metrics: list[dict[str, object]] = []

    changes = inventory_changes()
    changes.to_csv(OUT / "inventory_change.csv", index=False, encoding="utf-8-sig")
    for state in ("added", "changed", "removed", "unchanged"):
        metric(metrics, "파일 변경", state, int((changes["change"] == state).sum()))

    # Administrative-dong boundary and code workbook.
    adm_dbf_path = find_path("BND_ADM_DONG_PG", ".dbf")
    adm_shp_path = find_path("BND_ADM_DONG_PG", ".shp")
    adm = core.parse_dbf(adm_dbf_path, "cp949")
    adm_shp = profiler.read_shp_header(adm_shp_path)
    dalseo_adm = adm[adm["ADM_CD"].str.startswith("22070")].copy()
    population = core.read_csv(core.find_one("노인인구"))
    population_names = set(population["동명"].map(core.norm_name))
    boundary_names = set(dalseo_adm["ADM_NM"].map(core.norm_name))
    boundary_name_matches = population_names & boundary_names

    xlsx_path = find_path("센서스 공간정보 지역 코드", ".xlsx")
    workbook_dalseo, workbook_profile = latest_code_sheet_profile(xlsx_path)
    workbook_codes = set(workbook_dalseo["ADM_CD"])
    dbf_codes = set(dalseo_adm["ADM_CD"])
    dalseo_adm.sort_values("ADM_CD").to_csv(OUT / "dalseo_admin_boundary_attributes.csv", index=False, encoding="utf-8-sig")
    workbook_dalseo.sort_values("ADM_CD").to_csv(OUT / "dalseo_admin_code_workbook_rows.csv", index=False, encoding="utf-8-sig")

    metric(metrics, "행정동 경계", "shp_records", adm_shp["record_count"])
    metric(metrics, "행정동 경계", "dbf_records", len(adm))
    metric(metrics, "행정동 경계", "dalseo_records", len(dalseo_adm))
    metric(metrics, "행정동 경계", "base_date", dalseo_adm["BASE_DATE"].iloc[0])
    metric(metrics, "행정동 경계", "population_name_matches", len(boundary_name_matches), f"denominator={len(population_names)}")
    metric(metrics, "행정동 경계", "duplicate_adm_codes", int(dalseo_adm.duplicated("ADM_CD", keep=False).sum()))
    metric(metrics, "행정동 코드 워크북", "sheet_count", workbook_profile["workbook_sheets"])
    metric(metrics, "행정동 코드 워크북", "latest_sheet_data_rows", workbook_profile["latest_sheet_data_rows"])
    metric(metrics, "행정동 코드 워크북", "dalseo_rows", workbook_profile["dalseo_rows"])
    metric(metrics, "행정동 코드 워크북", "dbf_code_matches", len(workbook_codes & dbf_codes), f"denominator={len(dbf_codes)}")

    # Newly added CSV sources.
    protection = read_source_csv("노인보호구역")
    protection["CCTV설치대수_num"] = pd.to_numeric(protection["CCTV설치대수"], errors="coerce")
    cctv_contradictions = int(((protection["CCTV설치여부"] == "Y") & (protection["CCTV설치대수_num"] <= 0)).sum())
    metric(metrics, "노인보호구역", "rows", len(protection))
    metric(metrics, "노인보호구역", "unique_coordinates", protection[["위도", "경도"]].drop_duplicates().shape[0])
    metric(metrics, "노인보호구역", "cctv_yes_rows", int((protection["CCTV설치여부"] == "Y").sum()))
    metric(metrics, "노인보호구역", "cctv_positive_count_rows", int((protection["CCTV설치대수_num"] > 0).sum()))
    metric(metrics, "노인보호구역", "cctv_cross_field_contradictions", cctv_contradictions)
    metric(metrics, "노인보호구역", "data_date", protection["데이터기준일자"].iloc[0])

    ped = read_source_csv("보행자전용도로_20260707")
    segment_columns = [
        "보행자전용도로시작점위도", "보행자전용도로시작점경도",
        "보행자전용도로종료점위도", "보행자전용도로종료점경도",
    ]
    safety_columns = [
        "유지보수내용", "영상정보기처리기기설치개수", "보안등설치개수", "횡단보도설치개수",
        "방호울타리설치개수", "차량진입억제용말뚝설치개수", "속도저감시설설치개수",
        "교통표지판설치개수", "이정표설치개수", "점자블럭설치개수",
    ]
    all_blank_safety_columns = sum(int((ped[column] == "").all()) for column in safety_columns)
    metric(metrics, "달서구 보행자전용도로 CSV", "rows", len(ped))
    metric(metrics, "달서구 보행자전용도로 CSV", "unique_road_names", ped["보행자전용도로명"].nunique())
    metric(metrics, "달서구 보행자전용도로 CSV", "unique_segments", ped[segment_columns].drop_duplicates().shape[0])
    metric(metrics, "달서구 보행자전용도로 CSV", "blank_designation_date_rows", int((ped["지정일자"] == "").sum()))
    metric(metrics, "달서구 보행자전용도로 CSV", "all_blank_safety_columns", all_blank_safety_columns, " | ".join(safety_columns))
    metric(metrics, "달서구 보행자전용도로 CSV", "management_inspection_date", ped["관리점검일자"].iloc[0])
    metric(metrics, "달서구 보행자전용도로 CSV", "data_date", ped["데이터기준일자"].iloc[0])

    bridge = read_source_csv("육교정보")
    bridge_coordinate_duplicates = int(bridge.duplicated(["위도", "경도"], keep=False).sum())
    metric(metrics, "달서구 육교", "rows", len(bridge))
    metric(metrics, "달서구 육교", "unique_coordinates", bridge[["위도", "경도"]].drop_duplicates().shape[0])
    metric(metrics, "달서구 육교", "duplicate_coordinate_rows", bridge_coordinate_duplicates)
    metric(metrics, "달서구 육교", "accessible_yes_rows", int((bridge["장애인편의시설설치여부"] == "Y").sum()))
    metric(metrics, "달서구 육교", "elevator_rows", int(bridge["부대시설종류"].str.contains("엘리베이터", na=False).sum()))
    metric(metrics, "달서구 육교", "data_date", bridge["데이터기준일자"].iloc[0])

    # New spatial safety/walking inventories.
    spatial_rows = []
    for token, label, encoding in [
        ("A0043325", "횡단보도", "utf-8"),
        ("C_UQ164_PW", "공공보행통로", "utf-8"),
        ("C_UQ164_RF", "보행자우선도로", "utf-8"),
        ("C_UQ164_RO", "보행자전용도로 공간정보", "utf-8"),
    ]:
        dbf_path = find_path(token, ".dbf")
        shp_path = find_path(token, ".shp")
        dbf = core.parse_dbf(dbf_path, encoding)
        shp = profiler.read_shp_header(shp_path)
        daegu_rows = int((dbf["SIGNGU_SE"] == "27000").sum()) if "SIGNGU_SE" in dbf.columns else ""
        spatial_rows.append(
            {
                "dataset": label,
                "shape_type": shp["shape_type"],
                "shp_records": shp["record_count"],
                "dbf_records": len(dbf),
                "attributes": " | ".join(dbf.columns),
                "daegu_rows_by_SIGNGU_SE_27000": daegu_rows,
                "has_dalseo_attribute": False,
            }
        )
    spatial = pd.DataFrame(spatial_rows)
    spatial.to_csv(OUT / "updated_spatial_layer_inventory.csv", index=False, encoding="utf-8-sig")
    for _, row in spatial.iterrows():
        metric(metrics, row["dataset"], "shp_records", row["shp_records"])
        metric(metrics, row["dataset"], "dbf_records", row["dbf_records"])
        if row["daegu_rows_by_SIGNGU_SE_27000"] != "":
            metric(metrics, row["dataset"], "daegu_rows_by_city_code_27000", row["daegu_rows_by_SIGNGU_SE_27000"])

    updated_findings = [
        {
            "finding_id": "DQ-10", "severity": "High", "confidence": "High",
            "dataset": "보행 관련 신규 자료",
            "finding": "보행자전용도로·보행자우선도로·공공보행통로 자료는 지정 구간 또는 도시계획 면형 자료이며 완전한 연결형 보행망이 아님",
            "impact": "이를 보행망으로 대체하면 일반 보도·골목이 누락되어 최단경로와 300m·500m 공백이 왜곡됨",
            "evidence": f"달서구 CSV {ped[segment_columns].drop_duplicates().shape[0]}개 고유 구간; 공간 레이어는 polygon(shape_type=5)",
            "remediation": "OSMnx walk 네트워크를 기본망으로 생성하고 신규 레이어는 보완·검증 속성으로 사용",
        },
        {
            "finding_id": "DQ-11", "severity": "Medium", "confidence": "High",
            "dataset": "달서구 보행자전용도로 CSV",
            "finding": "7행이 3개 고유 구간의 운영시간 행으로 반복되고 안전시설 관련 10개 컬럼이 전부 비어 있음",
            "impact": "행 수를 도로 수로 세면 중복집계되고 설치환경·안전성을 평가할 수 없음",
            "evidence": f"7행, 고유 구간 3개, 지정일자 공백 {(ped['지정일자'] == '').sum()}행, 관리점검일 2020-05-30",
            "remediation": "구간 ID와 운영시간표를 분리하고 안전시설은 다른 자료 또는 현장조사로 보완",
        },
        {
            "finding_id": "DQ-12", "severity": "Medium", "confidence": "High",
            "dataset": "횡단보도 공간정보",
            "finding": "대구 전체 17,573개 면형 객체이나 행정구·신호기·턱낮춤 등 안전속성이 없음",
            "impact": "달서구 필터와 고령친화 안전성 평가를 속성만으로 수행할 수 없음",
            "evidence": "속성: g2_id, uqftid_id, cnstc_se, cnstc_de; 기준 2023-08-16",
            "remediation": "행정동 경계와 공간조인하고 전국횡단보도표준데이터의 안전속성으로 보완",
        },
        {
            "finding_id": "DQ-13", "severity": "Medium", "confidence": "High",
            "dataset": "노인보호구역",
            "finding": "자료가 2022년 5개 구역에 한정되고 CCTV 설치 여부와 설치대수가 모순인 행이 1개 있음",
            "impact": "현재 안전상태의 완전한 대리변수로 사용할 수 없음",
            "evidence": f"5행, 기준일 2022-10-07, CCTV Y이면서 0대 {cctv_contradictions}행",
            "remediation": "최신 원장 확인 후 구역 존재 여부 중심의 보조지표로만 사용",
        },
        {
            "finding_id": "DQ-14", "severity": "Medium", "confidence": "High",
            "dataset": "도시계획 보행 공간정보",
            "finding": "시군구 속성이 27000(대구 전체)까지만 제공되어 달서구를 속성 필터할 수 없음",
            "impact": "그대로 사용하면 대구 전역 객체가 달서구 분석에 섞임",
            "evidence": "대구 코드 객체: 공공보행통로 15, 보행자우선도로 1, 보행자전용도로 9",
            "remediation": "서로 다른 좌표계를 통일한 뒤 달서구 행정경계와 공간조인",
        },
        {
            "finding_id": "DQ-15", "severity": "Low", "confidence": "Medium",
            "dataset": "달서구 육교",
            "finding": "10개 육교 중 2개 명칭이 동일 좌표를 공유함",
            "impact": "실제 동일 교차지점의 별도 시설인지 좌표 중복인지 확인 전 자동 병합·삭제하면 안 됨",
            "evidence": f"10행, 고유 좌표 9개, 중복좌표 관련 {bridge_coordinate_duplicates}행",
            "remediation": "주소·현장지도를 확인하고 계단·엘리베이터 연결정보를 별도 구축",
        },
        {
            "finding_id": "DQ-16", "severity": "Medium", "confidence": "High",
            "dataset": "행정동 코드 워크북",
            "finding": "28개 연도 시트에서 코드 셀의 자료형·길이 표현이 달라질 수 있음",
            "impact": "숫자로 읽으면 선행 0이 사라져 경계키 조인이 깨질 수 있음",
            "evidence": "최신 2025년 6월 시트에서 달서구 23개 코드가 DBF 23개와 전부 일치",
            "remediation": "2025년 6월 시트를 기준으로 각 코드 조각을 문자열 zfill 후 결합",
        },
    ]
    previous_findings = pd.read_csv(OUT / "quality_findings.csv")
    previous_findings["confidence"] = "High"
    previous_findings["remediation"] = "기존 데이터 품질 점검표의 처리 원칙 적용"
    combined = pd.concat([previous_findings, pd.DataFrame(updated_findings)], ignore_index=True, sort=False)
    combined.to_csv(OUT / "updated_quality_findings.csv", index=False, encoding="utf-8-sig")

    metrics_frame = pd.DataFrame(metrics)
    metrics_frame.to_csv(OUT / "updated_source_metrics.csv", index=False, encoding="utf-8-sig")
    summary = {
        "csv_files_current": int(pd.read_csv(OUT / "file_inventory.csv").shape[0]),
        "csv_rows_current": int(pd.read_csv(OUT / "file_inventory.csv")["rows"].sum()),
        "spatial_layers_current": int(pd.read_csv(OUT / "spatial_profile.csv").query("extension == '.shp'").shape[0]),
        "new_csv_files": int((changes["change"] == "added").sum()),
        "changed_existing_csv_files": int((changes["change"] == "changed").sum()),
        "admin_boundary_ready": bool(len(dalseo_adm) == 23 and len(boundary_name_matches) == 23 and len(workbook_codes & dbf_codes) == 23),
        "complete_walk_network_ready": False,
        "quality_findings_total": int(len(combined)),
        "quality_findings_by_severity": combined["severity"].value_counts().to_dict(),
    }
    (OUT / "updated_validation_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False))


if __name__ == "__main__":
    main()
